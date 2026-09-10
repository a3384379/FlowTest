import json
import shlex
from base64 import b64encode
from dataclasses import dataclass
from datetime import UTC
from enum import StrEnum
from io import BytesIO
from typing import cast
from urllib.parse import urlencode
from uuid import UUID, uuid4

from openpyxl import Workbook
from pydantic import JsonValue
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError
from app.core.logging import redact
from app.models.access import User
from app.models.api_assets import APIDefinition, APIVersion
from app.repositories.api_assets import APIAssetRepository
from app.services.projects import ProjectService


class APIExportFormat(StrEnum):
    HAR = "har"
    CURL = "curl"
    BRUNO = "bruno"
    EXCEL = "excel"


@dataclass(frozen=True, slots=True)
class APIExportDocument:
    filename: str
    media_type: str
    content: bytes


@dataclass(frozen=True, slots=True)
class ExportedAPI:
    definition: APIDefinition
    version: APIVersion


class APIExportService:
    def __init__(self, session: AsyncSession) -> None:
        self._assets = APIAssetRepository(session)
        self._projects = ProjectService(session)

    async def export(
        self, *, actor: User, project_id: UUID, export_format: APIExportFormat
    ) -> APIExportDocument:
        access = await self._projects.get(actor=actor, project_id=project_id)
        definitions, total = await self._assets.list_definitions(
            project_id=project_id,
            offset=0,
            limit=10_000,
        )
        if total > 10_000:
            raise AppError(
                code="EXPORT_LIMIT_EXCEEDED", message="接口数量超过导出上限 10000", status_code=422
            )
        items = await self._load_versions(definitions)
        safe_name = (
            "".join(
                character if character.isalnum() or character in {"-", "_"} else "-"
                for character in access.project.name
            ).strip("-")
            or "flowtest"
        )
        if export_format is APIExportFormat.HAR:
            return APIExportDocument(
                f"{safe_name}.har",
                "application/json",
                _json_bytes(_har_document(items)),
            )
        if export_format is APIExportFormat.CURL:
            commands = (
                "# 地址为模板; 请将 flowtest.invalid 替换为明确的环境或服务地址。\n"
                + "\n\n".join(_curl_command(item) for item in items)
            )
            return APIExportDocument(f"{safe_name}.curl.txt", "text/plain", commands.encode())
        if export_format is APIExportFormat.BRUNO:
            return APIExportDocument(
                f"{safe_name}.bruno.json",
                "application/json",
                _json_bytes(_bruno_document(access.project.name, items)),
            )
        return APIExportDocument(
            f"{safe_name}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _excel_document(items),
        )

    async def _load_versions(self, definitions: list[APIDefinition]) -> list[ExportedAPI]:
        items: list[ExportedAPI] = []
        for definition in definitions:
            version = await self._assets.get_version(
                definition_id=definition.id,
                version=definition.current_version,
            )
            if version is None:
                raise AppError(
                    code="API_VERSION_NOT_FOUND",
                    message="导出接口的当前版本不存在",
                    status_code=409,
                    details={"api_id": str(definition.id), "version": definition.current_version},
                )
            items.append(ExportedAPI(definition, version))
        return items


def _har_document(items: list[ExportedAPI]) -> dict[str, JsonValue]:
    return {
        "log": {
            "version": "1.2",
            "creator": {"name": "FlowTest", "version": "2"},
            "entries": [
                {
                    "comment": item.definition.name,
                    "request": _har_request(item.version),
                    "response": {
                        "status": 0,
                        "statusText": "",
                        "httpVersion": "HTTP/1.1",
                        "headers": [],
                        "cookies": [],
                        "content": {"size": 0, "mimeType": "application/json"},
                        "redirectURL": "",
                        "headersSize": -1,
                        "bodySize": -1,
                    },
                    "startedDateTime": item.version.created_at.replace(
                        tzinfo=item.version.created_at.tzinfo or UTC
                    ).isoformat(),
                    "time": 0,
                    "cache": {},
                    "timings": {"send": 0, "wait": 0, "receive": 0},
                }
                for item in items
            ],
        }
    }


def _curl_command(item: ExportedAPI) -> str:
    parts = ["curl", "-X", item.version.method, shlex.quote(_export_url(item.version))]
    for name, value in _export_headers(item.version).items():
        parts.extend(["-H", shlex.quote(f"{name}: {value}")])
    if item.version.body_kind == "multipart":
        for name, value in _multipart_fields(item.version).items():
            parts.extend(["--form-string", shlex.quote(f"{name}={value}")])
    elif item.version.body_kind != "none":
        parts.extend(["--data-raw", shlex.quote(_body_text(item.version))])
    return " ".join(parts)


def _bruno_document(project_name: str, items: list[ExportedAPI]) -> dict[str, JsonValue]:
    return {
        "version": "1",
        "uid": _bruno_uid(),
        "name": project_name,
        "items": [
            {
                "uid": _bruno_uid(),
                "name": item.definition.name,
                "description": item.definition.description,
                "type": "http-request",
                "seq": index,
                "request": {
                    "method": item.version.method,
                    "url": _export_url(item.version, include_auth=False),
                    "headers": [
                        {"uid": _bruno_uid(), "name": name, "value": value, "enabled": True}
                        for name, value in _bruno_headers(item.version).items()
                    ],
                    "params": [],
                    "auth": _bruno_auth(item.version),
                    "body": _bruno_body(item.version),
                },
            }
            for index, item in enumerate(items, 1)
        ],
    }


def _excel_document(items: list[ExportedAPI]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Excel workbook has no active sheet")
    sheet.title = "APIs"
    sheet.append(
        [
            "name",
            "method",
            "path",
            "description",
            "query",
            "headers",
            "body",
            "auth_kind",
            "auth_config",
            "body_kind",
        ]
    )
    for item in items:
        query = [
            {"name": name, "value": value, "enabled": True}
            for name, value in _safe_query_items(item.version)
        ]
        sheet.append(
            [
                _excel_safe_cell(item.definition.name),
                _excel_safe_cell(item.version.method),
                _excel_safe_cell(item.version.path),
                _excel_safe_cell(item.definition.description),
                _excel_safe_cell(json.dumps(query, ensure_ascii=False)),
                _excel_safe_cell(
                    json.dumps(_safe_headers(item.version.headers), ensure_ascii=False)
                ),
                _excel_safe_cell(json.dumps(_safe_body(item.version), ensure_ascii=False))
                if item.version.body is not None
                else "",
                _excel_safe_cell(item.version.auth_kind),
                _excel_safe_cell(
                    json.dumps(_safe_auth(item.version.auth_config), ensure_ascii=False)
                ),
                item.version.body_kind,
            ]
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _request_target(version: APIVersion, *, include_auth: bool = True) -> str:
    query = _safe_query_items(version) + (_auth_query(version) if include_auth else [])
    return f"{version.path}?{urlencode(query)}" if query else version.path


def _har_post_data(version: APIVersion) -> JsonValue:
    if version.body_kind == "multipart":
        return {
            "mimeType": "multipart/form-data",
            "params": [
                {"name": name, "value": value} for name, value in _multipart_fields(version).items()
            ],
        }
    return {"mimeType": _body_mime(version), "text": _body_text(version)}


def _har_request(version: APIVersion) -> dict[str, JsonValue]:
    request: dict[str, JsonValue] = {
        "method": version.method,
        "url": _export_url(version),
        "httpVersion": "HTTP/1.1",
        "cookies": [],
        "headersSize": -1,
        "bodySize": -1,
        "headers": [
            {"name": name, "value": value} for name, value in _export_headers(version).items()
        ],
        "queryString": [
            {"name": name, "value": value}
            for name, value in _safe_query_items(version) + _auth_query(version)
        ],
        "comment": "地址为模板, 请替换 flowtest.invalid 为明确的环境或服务地址",
    }
    if version.body_kind != "none":
        request["postData"] = _har_post_data(version)
    return request


def _export_url(version: APIVersion, *, include_auth: bool = True) -> str:
    path = _request_target(version, include_auth=include_auth)
    if path.startswith(("http://", "https://")):
        return path
    return f"http://flowtest.invalid/{path.lstrip('/')}"


def _export_headers(version: APIVersion) -> dict[str, str]:
    headers = _safe_headers({**version.headers, **_export_auth_headers(version)})
    if version.body_kind in {"json", "form", "raw"} and not any(
        name.lower() == "content-type" for name in headers
    ):
        headers = {**headers, "Content-Type": _body_mime(version)}
    return headers


def _multipart_fields(version: APIVersion) -> dict[str, str]:
    body = _safe_body(version)
    if not isinstance(body, dict):
        return {}
    if body.get("files"):
        raise AppError(
            code="EXPORT_MULTIPART_FILES_UNSUPPORTED",
            message="该导出格式无法携带上传文件, 请使用 Excel 保留文件引用并在目标项目重新绑定",
            status_code=422,
        )
    fields = body.get("fields", {})
    return {name: str(value) for name, value in fields.items()} if isinstance(fields, dict) else {}


def _safe_query_items(version: APIVersion) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    for parameter in version.query_parameters:
        if not parameter.get("enabled", True):
            continue
        name = str(parameter["name"])
        safe = cast(dict[str, str], redact({name: str(parameter.get("value", ""))}))
        result.append((name, safe[name]))
    return result


def _safe_body(version: APIVersion) -> JsonValue:
    return cast(JsonValue, redact(version.body))


def _safe_headers(headers: dict[str, str]) -> dict[str, str]:
    return cast(dict[str, str], redact(headers))


def _safe_auth(auth: dict[str, str]) -> dict[str, str]:
    return cast(dict[str, str], redact(auth))


def _excel_safe_cell(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _json_bytes(value: dict[str, JsonValue]) -> bytes:
    return json.dumps(value, ensure_ascii=False, indent=2).encode()


def _body_mime(version: APIVersion) -> str:
    return {
        "json": "application/json",
        "form": "application/x-www-form-urlencoded",
        "multipart": "multipart/form-data",
    }.get(version.body_kind, "text/plain")


def _body_text(version: APIVersion) -> str:
    body = _safe_body(version)
    if version.body_kind == "raw" and isinstance(body, str):
        return body
    if version.body_kind == "form" and isinstance(body, dict):
        return urlencode(body)
    return json.dumps(body, ensure_ascii=False)


def _bruno_body(version: APIVersion) -> dict[str, JsonValue]:
    body = _safe_body(version)
    if version.body_kind == "none":
        return {"mode": "none"}
    if version.body_kind == "form" and isinstance(body, dict):
        return {
            "mode": "formUrlEncoded",
            "formUrlEncoded": [
                {"uid": _bruno_uid(), "name": name, "value": str(value), "enabled": True}
                for name, value in body.items()
            ],
        }
    if version.body_kind == "raw":
        return {"mode": "text", "text": _body_text(version)}
    if version.body_kind == "multipart":
        return {
            "mode": "multipartForm",
            "multipartForm": [
                {"uid": _bruno_uid(), "name": name, "value": value, "type": "text", "enabled": True}
                for name, value in _multipart_fields(version).items()
            ],
        }
    return {"mode": "json", "json": _body_text(version)}


def _bruno_auth(version: APIVersion) -> dict[str, JsonValue]:
    auth = _safe_auth(version.auth_config)
    if version.auth_kind in {"bearer", "basic"}:
        return {"mode": version.auth_kind, version.auth_kind: cast(JsonValue, auth)}
    if version.auth_kind == "api_key" and auth.get("in") == "cookie":
        return {"mode": "none"}
    if version.auth_kind == "api_key":
        return {
            "mode": "apikey",
            "apikey": {
                "key": auth.get("name", ""),
                "value": auth.get("value", ""),
                "placement": "queryparams" if auth.get("in") == "query" else "header",
            },
        }
    return {"mode": "none"}


def _export_auth_headers(version: APIVersion) -> dict[str, str]:
    auth = version.auth_config
    if version.auth_kind == "bearer":
        return {"Authorization": f"Bearer {auth.get('token', '')}"}
    if version.auth_kind == "basic":
        value = b64encode(
            f"{auth.get('username', '')}:{auth.get('password', '')}".encode()
        ).decode()
        return {"Authorization": f"Basic {value}"}
    if version.auth_kind == "api_key" and auth.get("in", "header") == "header":
        return {auth.get("name", "X-API-Key"): auth.get("value", "")}
    if version.auth_kind == "api_key" and auth.get("in") == "cookie":
        value = f"{auth.get('name', 'api_key')}={auth.get('value', '')}"
        previous = version.headers.get("Cookie", "")
        return {"Cookie": f"{previous}; {value}" if previous else value}
    return {}


def _auth_query(version: APIVersion) -> list[tuple[str, str]]:
    if version.auth_kind != "api_key" or version.auth_config.get("in") != "query":
        return []
    auth = _safe_auth(version.auth_config)
    return [(auth.get("name", "X-API-Key"), auth.get("value", ""))]


def _bruno_uid() -> str:
    return uuid4().hex[:21]


def _bruno_headers(version: APIVersion) -> dict[str, str]:
    if version.auth_kind == "api_key" and version.auth_config.get("in") == "cookie":
        return _export_headers(version)
    return _safe_headers(version.headers)
