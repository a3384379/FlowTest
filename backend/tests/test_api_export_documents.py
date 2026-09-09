import json
import shlex
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.core.errors import AppError
from app.models.api_assets import APIDefinition, APIVersion
from app.services.api_exports import (
    APIExportFormat,
    APIExportService,
    ExportedAPI,
    _bruno_document,
    _curl_command,
    _excel_document,
    _har_document,
)


def api_item(body_kind: str = "json", body: object = None) -> ExportedAPI:
    definition = APIDefinition(id=uuid4(), name="中文接口", description="说明", current_version=1)
    version = APIVersion(
        method="POST",
        path="/orders",
        headers={},
        body_kind=body_kind,
        body=body,
        auth_kind="none",
        auth_config={},
        created_at=datetime.now(UTC),
        query_parameters=[
            {"name": "tag", "value": "a"},
            {"name": "tag", "value": "b"},
            {"name": "skip", "value": "x", "enabled": False},
        ],
    )
    return ExportedAPI(definition, version)


def test_bruno_export_has_native_request_shape() -> None:
    document = _bruno_document(
        "中文项目", [api_item("json", {"zero": 0, "false": False, "null": None})]
    )
    assert document["version"] == "1"
    item = document["items"][0]
    assert item["type"] == "http-request"
    assert len(document["uid"]) == len(item["uid"]) == 21
    assert isinstance(item["request"]["headers"], list)
    assert item["request"]["body"]["mode"] == "json"
    assert json.loads(item["request"]["body"]["json"]) == {"zero": 0, "false": False, "null": None}


@pytest.mark.parametrize(
    "kind,body,expected",
    [
        ("raw", "hello", "hello"),
        ("form", {"name": "中文", "count": 0}, "name=%E4%B8%AD%E6%96%87&count=0"),
    ],
)
def test_curl_keeps_body_encoding(kind: str, body: object, expected: str) -> None:
    command = shlex.split(_curl_command(api_item(kind, body)))
    assert command[command.index("--data-raw") + 1] == expected
    assert "tag=a&tag=b" in command[3]


def test_har_required_fields_and_excel_duplicate_query() -> None:
    item = api_item("json", False)
    entry = _har_document([item])["log"]["entries"][0]
    assert entry["timings"] == {"send": 0, "wait": 0, "receive": 0}
    assert entry["request"]["cookies"] == []
    assert entry["request"]["postData"]["text"] == "false"
    sheet = load_workbook(BytesIO(_excel_document([item]))).active
    assert json.loads(sheet.cell(2, 5).value) == [
        {"name": "tag", "value": "a", "enabled": True},
        {"name": "tag", "value": "b", "enabled": True},
    ]


@pytest.mark.asyncio
async def test_export_refuses_truncation_and_missing_versions() -> None:
    service = APIExportService(AsyncMock())
    service._projects = SimpleNamespace(
        get=AsyncMock(return_value=SimpleNamespace(project=SimpleNamespace(name="项目")))
    )
    service._assets = SimpleNamespace(list_definitions=AsyncMock(return_value=([], 10001)))
    with pytest.raises(AppError, match="导出上限"):
        await service.export(
            actor=SimpleNamespace(), project_id=uuid4(), export_format=APIExportFormat.HAR
        )
    item = api_item()
    service._assets = SimpleNamespace(
        list_definitions=AsyncMock(return_value=([item.definition], 1)),
        get_version=AsyncMock(return_value=None),
    )
    with pytest.raises(AppError, match="版本不存在"):
        await service.export(
            actor=SimpleNamespace(), project_id=uuid4(), export_format=APIExportFormat.HAR
        )


@pytest.mark.parametrize(
    "kind,body",
    [
        ("none", None),
        ("json", None),
        ("json", 0),
        ("json", False),
        ("raw", "raw text"),
        ("form", {"field": "value"}),
        ("multipart", {"fields": {"field": "value"}, "files": []}),
    ],
)
def test_exported_documents_can_be_imported_without_body_or_query_loss(
    kind: str, body: object
) -> None:
    from app.importers.excel import parse_excel
    from app.importers.http_formats import parse_bruno

    item = api_item(kind, body)
    excel = parse_excel(_excel_document([item]))[0].request
    document = _bruno_document("项目", [item])
    bruno = parse_bruno(json.dumps(document).encode(), document)[0].request
    for request in (excel, bruno):
        assert request.body_kind.value == kind
        assert request.body == body
        assert [(value.name, value.value) for value in request.query_parameters] == [
            ("tag", "a"),
            ("tag", "b"),
        ]


def test_multipart_fields_export_and_file_references_fail_explicitly() -> None:
    item = api_item("multipart", {"fields": {"name": "中文"}, "files": []})
    assert "--form-string" in _curl_command(item)
    assert _bruno_document("项目", [item])["items"][0]["request"]["body"]["mode"] == "multipartForm"
    assert _har_document([item])["log"]["entries"][0]["request"]["postData"]["params"] == [
        {"name": "name", "value": "中文"}
    ]
    with pytest.raises(AppError, match="上传文件"):
        _curl_command(
            api_item(
                "multipart",
                {"fields": {}, "files": [{"field": "file", "artifact_id": str(uuid4())}]},
            )
        )


def test_har_sqlite_timestamp_has_timezone() -> None:
    item = api_item()
    item.version.created_at = datetime(2026, 9, 9, 12, 0)
    assert _har_document([item])["log"]["entries"][0]["startedDateTime"].endswith("+00:00")
